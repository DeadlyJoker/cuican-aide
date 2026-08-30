"""Slide type renderers for build_pptx.py.

Each entry in a deck spec's ``slides`` array picks one renderer by ``layout``.
Renderers only lay out geometry; colors and type scale come from the Theme.
"""

from __future__ import annotations

from typing import Any, Callable

from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

from . import spec as specmod
from .pptxlayout import (
    Theme,
    add_accent_bar,
    add_card,
    add_picture_fit,
    add_text,
    blank_slide,
    paint_background,
)

MARGIN = 0.7
CONTENT_TOP = 1.65


def _slide_size(presentation) -> tuple[float, float]:
    return presentation.slide_width / 914400, presentation.slide_height / 914400


def _title_block(slide, presentation, theme: Theme, block: dict[str, Any]) -> None:
    width, _ = _slide_size(presentation)
    add_accent_bar(slide, theme, left=MARGIN, top=0.62)
    add_text(
        slide,
        text=block.get("title", ""),
        left=MARGIN,
        top=0.8,
        width=width - MARGIN * 2,
        height=0.85,
        size=theme.heading_size,
        color=theme.rgb("ink"),
        font=theme.heading_font,
        bold=True,
    )
    subtitle = block.get("subtitle")
    if subtitle:
        add_text(
            slide,
            text=subtitle,
            left=MARGIN,
            top=1.55,
            width=width - MARGIN * 2,
            height=0.4,
            size=theme.caption_size + 1,
            color=theme.rgb("muted"),
            font=theme.body_font,
        )


def render_cover(slide, presentation, theme: Theme, block: dict[str, Any]) -> None:
    width, height = _slide_size(presentation)
    paint_background(slide, presentation, theme, "background")
    band = slide.shapes.add_shape(1, 0, 0, Inches(0.22), presentation.slide_height)
    band.fill.solid()
    band.fill.fore_color.rgb = theme.rgb("accent")
    band.line.fill.background()
    band.shadow.inherit = False

    add_text(
        slide,
        text=block.get("title", ""),
        left=1.1,
        top=height / 2 - 1.5,
        width=width - 2.2,
        height=1.6,
        size=theme.title_size,
        color=theme.rgb("ink"),
        font=theme.heading_font,
        bold=True,
    )
    subtitle = block.get("subtitle")
    if subtitle:
        add_text(
            slide,
            text=subtitle,
            left=1.1,
            top=height / 2 + 0.15,
            width=width - 2.2,
            height=0.9,
            size=theme.body_size + 2,
            color=theme.rgb("muted"),
            font=theme.body_font,
        )
    meta = block.get("meta")
    if meta:
        add_text(
            slide,
            text=meta,
            left=1.1,
            top=height - 1.25,
            width=width - 2.2,
            height=0.5,
            size=theme.caption_size,
            color=theme.rgb("muted"),
            font=theme.body_font,
        )


def render_section(slide, presentation, theme: Theme, block: dict[str, Any]) -> None:
    width, height = _slide_size(presentation)
    paint_background(slide, presentation, theme, "accent")
    add_text(
        slide,
        text=block.get("title", ""),
        left=MARGIN + 0.3,
        top=height / 2 - 0.8,
        width=width - (MARGIN + 0.3) * 2,
        height=1.2,
        size=theme.title_size - 6,
        color=theme.rgb("background"),
        font=theme.heading_font,
        bold=True,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    subtitle = block.get("subtitle")
    if subtitle:
        add_text(
            slide,
            text=subtitle,
            left=MARGIN + 0.3,
            top=height / 2 + 0.5,
            width=width - (MARGIN + 0.3) * 2,
            height=0.6,
            size=theme.body_size,
            color=theme.rgb("accent_soft"),
            font=theme.body_font,
        )


def render_bullets(slide, presentation, theme: Theme, block: dict[str, Any]) -> None:
    width, height = _slide_size(presentation)
    paint_background(slide, presentation, theme)
    _title_block(slide, presentation, theme, block)

    bullets = block.get("bullets") or []
    if not bullets:
        raise specmod.SpecError("bullets slide needs `bullets`")
    box = slide.shapes.add_textbox(
        Inches(MARGIN),
        Inches(CONTENT_TOP + 0.35),
        Inches(width - MARGIN * 2),
        Inches(height - CONTENT_TOP - 1.0),
    )
    frame = box.text_frame
    frame.word_wrap = True
    for index, item in enumerate(bullets):
        text, level = (item.get("text", ""), int(item.get("level", 0))) if isinstance(item, dict) else (item, 0)
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.line_spacing = 1.3
        paragraph.space_after = Pt(10)
        paragraph.level = min(level, 4)
        marker = paragraph.add_run()
        marker.text = "•  " if level == 0 else "–  "
        marker.font.size = Pt(theme.body_size - level)
        marker.font.color.rgb = theme.rgb("accent")
        marker.font.name = theme.body_font
        run = paragraph.add_run()
        run.text = str(text)
        run.font.size = Pt(theme.body_size - level)
        run.font.color.rgb = theme.rgb("ink" if level == 0 else "muted")
        run.font.name = theme.body_font


def render_cards(slide, presentation, theme: Theme, block: dict[str, Any]) -> None:
    width, height = _slide_size(presentation)
    paint_background(slide, presentation, theme)
    _title_block(slide, presentation, theme, block)

    cards = block.get("cards") or []
    if not cards:
        raise specmod.SpecError("cards slide needs `cards`")
    columns = int(block.get("columns", min(len(cards), 3)))
    rows = -(-len(cards) // columns)
    gap = 0.28
    card_w = (width - MARGIN * 2 - gap * (columns - 1)) / columns
    available = height - CONTENT_TOP - 0.9
    card_h = (available - gap * (rows - 1)) / rows

    for index, card in enumerate(cards):
        row, column = divmod(index, columns)
        left = MARGIN + column * (card_w + gap)
        top = CONTENT_TOP + 0.3 + row * (card_h + gap)
        add_card(slide, theme, left=left, top=top, width=card_w, height=card_h)
        add_text(
            slide,
            text=card.get("title", ""),
            left=left + 0.28,
            top=top + 0.26,
            width=card_w - 0.56,
            height=0.5,
            size=theme.body_size + 2,
            color=theme.rgb("ink"),
            font=theme.heading_font,
            bold=True,
        )
        body = card.get("body")
        if body:
            add_text(
                slide,
                text=body,
                left=left + 0.28,
                top=top + 0.88,
                width=card_w - 0.56,
                height=card_h - 1.15,
                size=theme.caption_size + 1.5,
                color=theme.rgb("muted"),
                font=theme.body_font,
                line_spacing=1.3,
            )


def render_metrics(slide, presentation, theme: Theme, block: dict[str, Any]) -> None:
    width, height = _slide_size(presentation)
    paint_background(slide, presentation, theme)
    _title_block(slide, presentation, theme, block)

    metrics = block.get("metrics") or []
    if not metrics:
        raise specmod.SpecError("metrics slide needs `metrics`")
    gap = 0.3
    card_w = (width - MARGIN * 2 - gap * (len(metrics) - 1)) / len(metrics)
    card_h = min(2.4, height - CONTENT_TOP - 1.4)
    top = CONTENT_TOP + 0.5

    for index, metric in enumerate(metrics):
        left = MARGIN + index * (card_w + gap)
        add_card(slide, theme, left=left, top=top, width=card_w, height=card_h, fill="accent_soft")
        add_text(
            slide,
            text=metric.get("value", ""),
            left=left + 0.2,
            top=top + 0.4,
            width=card_w - 0.4,
            height=0.9,
            size=theme.title_size - 6,
            color=theme.rgb("accent"),
            font=theme.heading_font,
            bold=True,
            align="center",
        )
        add_text(
            slide,
            text=metric.get("label", ""),
            left=left + 0.2,
            top=top + 1.35,
            width=card_w - 0.4,
            height=0.8,
            size=theme.caption_size + 1,
            color=theme.rgb("ink"),
            font=theme.body_font,
            align="center",
            line_spacing=1.25,
        )


def render_table(slide, presentation, theme: Theme, block: dict[str, Any]) -> None:
    width, height = _slide_size(presentation)
    paint_background(slide, presentation, theme)
    _title_block(slide, presentation, theme, block)

    header = block.get("header") or []
    rows = block.get("rows") or []
    if not header or not rows:
        raise specmod.SpecError("table slide needs `header` and `rows`")
    shape = slide.shapes.add_table(
        len(rows) + 1,
        len(header),
        Inches(MARGIN),
        Inches(CONTENT_TOP + 0.35),
        Inches(width - MARGIN * 2),
        Inches(min(0.42 * (len(rows) + 1), height - CONTENT_TOP - 1.1)),
    )
    table = shape.table
    for column, label in enumerate(header):
        cell = table.cell(0, column)
        cell.text = str(label)
        _style_cell(cell, theme, size=theme.caption_size + 1.5, bold=True, fill="accent", ink="background")
    for row_index, row in enumerate(rows, start=1):
        for column, value in enumerate(row):
            cell = table.cell(row_index, column)
            cell.text = "" if value is None else str(value)
            _style_cell(cell, theme, size=theme.caption_size + 1, bold=False, fill=None, ink="ink")


def _style_cell(cell, theme: Theme, *, size: float, bold: bool, fill: str | None, ink: str) -> None:
    cell.margin_left = Inches(0.12)
    cell.margin_right = Inches(0.12)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = theme.rgb(fill)
    for paragraph in cell.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = theme.rgb(ink)
            run.font.name = theme.body_font


def render_image(slide, presentation, theme: Theme, block: dict[str, Any]) -> None:
    width, height = _slide_size(presentation)
    paint_background(slide, presentation, theme)
    _title_block(slide, presentation, theme, block)
    image = block.get("image")
    if not image:
        raise specmod.SpecError("image slide needs `image`")
    caption = block.get("caption")
    box_h = height - CONTENT_TOP - (1.4 if caption else 0.9)
    add_picture_fit(
        slide, image, left=MARGIN, top=CONTENT_TOP + 0.3, width=width - MARGIN * 2, height=box_h
    )
    if caption:
        add_text(
            slide,
            text=caption,
            left=MARGIN,
            top=CONTENT_TOP + 0.3 + box_h + 0.12,
            width=width - MARGIN * 2,
            height=0.5,
            size=theme.caption_size,
            color=theme.rgb("muted"),
            font=theme.body_font,
            align="center",
        )


def render_two_column(slide, presentation, theme: Theme, block: dict[str, Any]) -> None:
    width, height = _slide_size(presentation)
    paint_background(slide, presentation, theme)
    _title_block(slide, presentation, theme, block)
    gap = 0.4
    column_w = (width - MARGIN * 2 - gap) / 2
    top = CONTENT_TOP + 0.35
    column_h = height - top - 0.8

    for index, column in enumerate(block.get("columns") or []):
        left = MARGIN + index * (column_w + gap)
        add_text(
            slide,
            text=column.get("title", ""),
            left=left,
            top=top,
            width=column_w,
            height=0.45,
            size=theme.body_size + 2,
            color=theme.rgb("accent"),
            font=theme.heading_font,
            bold=True,
        )
        add_text(
            slide,
            text="\n".join(column.get("body", [])) if isinstance(column.get("body"), list) else column.get("body", ""),
            left=left,
            top=top + 0.6,
            width=column_w,
            height=column_h - 0.6,
            size=theme.body_size - 1,
            color=theme.rgb("ink"),
            font=theme.body_font,
            line_spacing=1.35,
        )


def render_closing(slide, presentation, theme: Theme, block: dict[str, Any]) -> None:
    width, height = _slide_size(presentation)
    paint_background(slide, presentation, theme, "surface")
    add_text(
        slide,
        text=block.get("title", "Thank you"),
        left=MARGIN,
        top=height / 2 - 0.9,
        width=width - MARGIN * 2,
        height=1.1,
        size=theme.title_size - 4,
        color=theme.rgb("ink"),
        font=theme.heading_font,
        bold=True,
        align="center",
        anchor=MSO_ANCHOR.MIDDLE,
    )
    subtitle = block.get("subtitle")
    if subtitle:
        add_text(
            slide,
            text=subtitle,
            left=MARGIN,
            top=height / 2 + 0.3,
            width=width - MARGIN * 2,
            height=0.7,
            size=theme.body_size,
            color=theme.rgb("muted"),
            font=theme.body_font,
            align="center",
        )


RENDERERS: dict[str, Callable[..., None]] = {
    "cover": render_cover,
    "section": render_section,
    "bullets": render_bullets,
    "cards": render_cards,
    "metrics": render_metrics,
    "table": render_table,
    "image": render_image,
    "twoColumn": render_two_column,
    "closing": render_closing,
}


def render_slide(presentation, theme: Theme, block: dict[str, Any]):
    layout = block.get("layout", "bullets")
    renderer = RENDERERS.get(layout)
    if renderer is None:
        known = ", ".join(sorted(RENDERERS))
        raise specmod.SpecError(f"unknown slide layout `{layout}`; known layouts: {known}")
    slide = blank_slide(presentation)
    renderer(slide, presentation, theme, block)
    notes = block.get("notes")
    if notes:
        slide.notes_slide.notes_text_frame.text = str(notes)
    return slide

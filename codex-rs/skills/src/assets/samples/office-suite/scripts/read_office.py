#!/usr/bin/env python3
"""Read .docx / .pptx / .xlsx / .pdf into Markdown for inspection or editing.

This replaces the pandoc/markitdown dependency: everything runs on the same
Python packages the builders already need, so reading works on a clean machine.

Usage:
    python read_office.py report.docx
    python read_office.py deck.pptx --max-chars 20000
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from officelib import env  # noqa: E402

env.reexec_in_environment()

DEFAULT_MAX_CHARS = 40000


def read_docx(path: Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(path))
    lines: list[str] = []

    # Walk the body in document order so tables stay where the author put them.
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "").lower()
            if style.startswith("heading"):
                level = "".join(filter(str.isdigit, style)) or "1"
                lines.append(f"{'#' * min(int(level), 6)} {text}")
            elif "list bullet" in style:
                lines.append(f"- {text}")
            elif "list number" in style:
                lines.append(f"1. {text}")
            else:
                lines.append(text)
            lines.append("")
        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            lines.extend(_markdown_table(table))
            lines.append("")
    return "\n".join(lines)


def _markdown_table(table) -> list[str]:
    rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
    if not rows:
        return []
    width = len(rows[0])
    out = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    for row in rows[1:]:
        out.append("| " + " | ".join(row) + " |")
    return out


def read_pptx(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(str(path))
    lines: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        lines.append(f"## Slide {index}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    for line in text.split("\n"):
                        lines.append(f"- {line.strip()}" if line.strip() else "")
            elif shape.has_table:
                lines.extend(_markdown_pptx_table(shape.table))
            elif shape.shape_type == 13:  # PICTURE
                lines.append(f"- [image] {getattr(shape, 'name', 'picture')}")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"> notes: {notes}")
        lines.append("")
    return "\n".join(lines)


def _markdown_pptx_table(table) -> list[str]:
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    if not rows:
        return []
    out = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * len(rows[0])) + " |"]
    for row in rows[1:]:
        out.append("| " + " | ".join(row) + " |")
    return out


def read_xlsx(path: Path, *, formulas: bool) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), data_only=not formulas)
    lines: list[str] = []
    for worksheet in workbook.worksheets:
        lines.append(f"## Sheet: {worksheet.title} ({worksheet.max_row}x{worksheet.max_column})")
        for row in worksheet.iter_rows(values_only=True):
            if all(value is None for value in row):
                continue
            lines.append("| " + " | ".join("" if v is None else str(v) for v in row) + " |")
        lines.append("")
    return "\n".join(lines)


def read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    lines: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        lines.append(f"## Page {index}")
        lines.append((page.extract_text() or "").strip())
        lines.append("")
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Read an office document as Markdown")
    parser.add_argument("path")
    parser.add_argument("--max-chars", type=int, default=DEFAULT_MAX_CHARS)
    parser.add_argument(
        "--formulas",
        action="store_true",
        help="for .xlsx, show formulas instead of cached values",
    )
    args = parser.parse_args(argv)

    path = Path(args.path).expanduser()
    if not path.is_file():
        print(f"error: file not found: {path}", file=sys.stderr)
        return 2

    suffix = path.suffix.lower()
    readers = {
        ".docx": lambda: read_docx(path),
        ".dotx": lambda: read_docx(path),
        ".pptx": lambda: read_pptx(path),
        ".xlsx": lambda: read_xlsx(path, formulas=args.formulas),
        ".xlsm": lambda: read_xlsx(path, formulas=args.formulas),
        ".pdf": lambda: read_pdf(path),
    }
    reader = readers.get(suffix)
    if reader is None:
        print(
            f"error: unsupported extension `{suffix}`; supported: {', '.join(sorted(readers))}",
            file=sys.stderr,
        )
        return 2

    text = reader()
    if len(text) > args.max_chars:
        text = text[: args.max_chars] + f"\n\n[truncated at {args.max_chars} chars]"
    print(text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

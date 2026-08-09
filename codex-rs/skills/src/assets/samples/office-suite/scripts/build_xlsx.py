#!/usr/bin/env python3
"""Build a .xlsx workbook from a JSON spec.

Formulas are written as real Excel formulas (strings starting with '='), so the
file stays live for the user instead of being a frozen value dump.

Usage:
    python build_xlsx.py workbook.json --out model.xlsx
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from officelib import env  # noqa: E402

env.reexec_in_environment()

from openpyxl import Workbook  # noqa: E402
from openpyxl.chart import BarChart, LineChart, PieChart, Reference  # noqa: E402
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.table import Table, TableStyleInfo  # noqa: E402

from officelib import spec as specmod  # noqa: E402

THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CHARTS = {"bar": BarChart, "line": LineChart, "pie": PieChart}


def write_sheet(worksheet, sheet: dict, theme: dict) -> None:
    header_fill = PatternFill("solid", fgColor=theme.get("headerFill", "2F6FED").lstrip("#").upper())
    header_font = Font(bold=True, color=theme.get("headerInk", "FFFFFF").lstrip("#").upper(), size=11)

    header = sheet.get("header") or []
    start_row = int(sheet.get("startRow", 1))
    row_cursor = start_row

    title = sheet.get("title")
    if title:
        cell = worksheet.cell(row=row_cursor, column=1, value=title)
        cell.font = Font(bold=True, size=14)
        row_cursor += 2

    header_row = row_cursor
    if header:
        for column, label in enumerate(header, start=1):
            cell = worksheet.cell(row=header_row, column=column, value=label)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BORDER
        row_cursor += 1

    number_formats = sheet.get("numberFormats") or {}
    for row in sheet.get("rows") or []:
        for column, value in enumerate(row, start=1):
            cell = worksheet.cell(row=row_cursor, column=column, value=value)
            cell.border = BORDER
            letter = get_column_letter(column)
            if letter in number_formats:
                cell.number_format = number_formats[letter]
        row_cursor += 1
    last_row = row_cursor - 1

    for widths_key, setter in (("widths", "width"),):
        widths = sheet.get(widths_key) or {}
        for letter, value in widths.items():
            setattr(worksheet.column_dimensions[letter], setter, float(value))
    if not sheet.get("widths"):
        _autofit(worksheet, header, sheet.get("rows") or [])

    if header and sheet.get("freezeHeader", True):
        worksheet.freeze_panes = worksheet.cell(row=header_row + 1, column=1)

    if header and sheet.get("autoFilter", True) and last_row > header_row:
        span = f"A{header_row}:{get_column_letter(len(header))}{last_row}"
        if sheet.get("asTable"):
            table = Table(displayName=sheet.get("tableName", f"T_{worksheet.title}"), ref=span)
            table.tableStyleInfo = TableStyleInfo(
                name=sheet.get("tableStyle", "TableStyleMedium9"),
                showRowStripes=True,
            )
            worksheet.add_table(table)
        else:
            worksheet.auto_filter.ref = span

    _apply_conditional(worksheet, sheet.get("conditional") or [])
    _apply_charts(worksheet, sheet.get("charts") or [], header_row=header_row, last_row=last_row)


def _autofit(worksheet, header: list, rows: list) -> None:
    columns = max([len(header)] + [len(row) for row in rows] or [0])
    for column in range(1, columns + 1):
        longest = len(str(header[column - 1])) if column <= len(header) else 0
        for row in rows:
            if column <= len(row) and row[column - 1] is not None:
                longest = max(longest, len(str(row[column - 1])))
        worksheet.column_dimensions[get_column_letter(column)].width = min(max(longest + 4, 10), 48)


def _apply_conditional(worksheet, rules: list) -> None:
    for rule in rules:
        span = specmod.require(rule, "range")
        kind = rule.get("kind", "colorScale")
        if kind == "colorScale":
            worksheet.conditional_formatting.add(
                span,
                ColorScaleRule(
                    start_type="min",
                    start_color=rule.get("startColor", "F8696B").lstrip("#"),
                    end_type="max",
                    end_color=rule.get("endColor", "63BE7B").lstrip("#"),
                ),
            )
        elif kind == "cellIs":
            worksheet.conditional_formatting.add(
                span,
                CellIsRule(
                    operator=rule.get("operator", "greaterThan"),
                    formula=[str(rule.get("formula", "0"))],
                    fill=PatternFill("solid", fgColor=rule.get("fill", "FFC7CE").lstrip("#")),
                ),
            )
        else:
            raise specmod.SpecError(f"unknown conditional kind `{kind}`")


def _apply_charts(worksheet, charts: list, *, header_row: int, last_row: int) -> None:
    for chart_spec in charts:
        kind = chart_spec.get("kind", "bar")
        factory = CHARTS.get(kind)
        if factory is None:
            raise specmod.SpecError(f"unknown chart kind `{kind}`; use {', '.join(CHARTS)}")
        chart = factory()
        chart.title = chart_spec.get("title")
        data = Reference(
            worksheet,
            min_col=int(chart_spec.get("dataMinCol", 2)),
            max_col=int(chart_spec.get("dataMaxCol", 2)),
            min_row=header_row,
            max_row=int(chart_spec.get("dataMaxRow", last_row)),
        )
        categories = Reference(
            worksheet,
            min_col=int(chart_spec.get("catCol", 1)),
            min_row=header_row + 1,
            max_row=int(chart_spec.get("dataMaxRow", last_row)),
        )
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        chart.height = float(chart_spec.get("height", 8))
        chart.width = float(chart_spec.get("width", 16))
        worksheet.add_chart(chart, chart_spec.get("anchor", "H3"))


def build(spec: dict, out_path: Path) -> Path:
    sheets = spec.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        raise specmod.SpecError("`sheets` must be a non-empty list")

    workbook = Workbook()
    workbook.remove(workbook.active)
    theme = spec.get("theme", {})

    for index, sheet in enumerate(sheets):
        if not isinstance(sheet, dict):
            raise specmod.SpecError(f"sheets[{index}] must be an object")
        worksheet = workbook.create_sheet(title=str(sheet.get("name", f"Sheet{index + 1}"))[:31])
        try:
            write_sheet(worksheet, sheet, theme)
        except specmod.SpecError:
            raise
        except Exception as error:  # noqa: BLE001 - surface the failing sheet
            raise specmod.SpecError(f"sheets[{index}] failed: {error}") from error

    workbook.save(str(out_path))
    return out_path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build a .xlsx workbook from a JSON spec")
    parser.add_argument("spec", nargs="?", default="-", help="spec path, or - for stdin")
    parser.add_argument("--out", dest="out", default=None, help="output .xlsx path")
    args = parser.parse_args(argv)

    try:
        spec = specmod.load_spec(args.spec)
        out_path = specmod.resolve_output(spec, args.out, ".xlsx")
        build(spec, out_path)
    except specmod.SpecError as error:
        print(f"error: {error}", file=sys.stderr)
        return 2

    print(out_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
